"""Video/audio sources and Studio artifacts (chart / infographic / spreadsheet)."""
import io
import json
from pathlib import Path

import pytest

from app import db, store, studio
from app.config import ARTIFACTS_DIR
from app.rag import format_timestamp


def _upload(client, nb_id, name, content, mime="application/octet-stream"):
    return client.post(
        f"/api/notebooks/{nb_id}/sources/file",
        files={"file": (name, io.BytesIO(content), mime)},
    )


# ---------- video / audio sources ----------

def test_video_upload_transcribed_and_indexed(client, notebook):
    res = _upload(client, notebook["id"], "talk.mp4", b"fake-video-bytes", "video/mp4")
    assert res.status_code == 200, res.text
    src = res.json()
    assert src["kind"] == "video" and src["status"] == "ready"
    chunks = store.notebook_chunks(notebook["id"])
    assert len(chunks) == 1
    assert "wombat" in chunks[0]["text"] and "42 tugrik" in chunks[0]["text"]
    assert chunks[0]["page"] == 0  # transcript block start second


def test_audio_upload_kind(client, notebook):
    res = _upload(client, notebook["id"], "talk.mp3", b"fake-audio", "audio/mpeg")
    assert res.json()["kind"] == "audio"


def test_media_endpoint_serves_file(client, notebook):
    src = _upload(client, notebook["id"], "talk.mp4", b"fake-video-bytes", "video/mp4").json()
    res = client.get(f"/api/media/{src['id']}")
    assert res.status_code == 200
    assert res.headers["content-type"] == "video/mp4"
    assert res.content == b"fake-video-bytes"
    # range requests supported for playback seeking
    partial = client.get(f"/api/media/{src['id']}", headers={"Range": "bytes=0-3"})
    assert partial.status_code == 206 and partial.content == b"fake"
    assert client.get("/api/media/nonexistent").status_code == 404


def test_video_citation_carries_kind_and_timestamp(client, notebook):
    _upload(client, notebook["id"], "talk.mp4", b"fake-video-bytes", "video/mp4")
    res = client.post(f"/api/notebooks/{notebook['id']}/chat",
                      json={"question": "What do tickets cost?", "history": []})
    sources_line = [l for l in res.text.splitlines() if l.startswith("data: [")][0]
    payload = json.loads(sources_line.removeprefix("data: "))
    assert payload[0]["kind"] == "video"
    assert payload[0]["page"] == 0
    msgs = client.get(f"/api/notebooks/{notebook['id']}/messages").json()
    assert msgs[1]["citations"][0]["kind"] == "video"


def test_format_timestamp():
    assert format_timestamp(0) == "0:00"
    assert format_timestamp(75) == "1:15"
    assert format_timestamp(3671) == "1:01:11"


# ---------- spreadsheet ingestion ----------

def test_xlsx_upload_indexed(client, notebook):
    from openpyxl import Workbook
    wb = Workbook()
    ws = wb.active
    ws.title = "Data"
    ws.append(["City", "Population"])
    ws.append(["Ulan Bator", 1645000])
    buf = io.BytesIO()
    wb.save(buf)
    res = _upload(client, notebook["id"], "cities.xlsx", buf.getvalue())
    assert res.status_code == 200, res.text
    assert res.json()["kind"] == "sheet"
    chunks = store.notebook_chunks(notebook["id"])
    assert "Ulan Bator | 1645000" in chunks[0]["text"]
    assert "Sheet: Data" in chunks[0]["text"]


# ---------- studio artifacts ----------

def _seed(client, nb_id):
    _upload(client, nb_id, "facts.txt",
            b"The festival happens in March. Tickets cost 42 tugrik.", "text/plain")


@pytest.mark.parametrize("kind,title", [
    ("chart", "Mock Chart"),
    ("infographic", "Mock Infographic"),
    ("spreadsheet", "Mock Table"),
    ("mindgraph", "Mock Mind Graph"),
    ("briefing", "Mock Briefing"),
    ("study_guide", "Mock Study Guide"),
    ("faq", "Mock FAQ"),
    ("timeline", "Mock Timeline"),
    ("source_summary", "Mock Summary"),
])
def test_artifact_generation(client, notebook, kind, title):
    _seed(client, notebook["id"])
    res = client.post(f"/api/notebooks/{notebook['id']}/artifacts", json={"kind": kind})
    assert res.status_code == 200, res.text
    art = res.json()
    assert art["title"] == title and art["kind"] == kind
    assert art["spec"]["source_note"]

    listed = client.get(f"/api/notebooks/{notebook['id']}/artifacts").json()
    assert len(listed) == 1

    if kind == "spreadsheet":
        assert art["file_url"]
        f = client.get(art["file_url"])
        assert f.status_code == 200
        assert "spreadsheetml" in f.headers["content-type"]
        from openpyxl import load_workbook
        wb = load_workbook(io.BytesIO(f.content))
        rows = list(wb.active.iter_rows(values_only=True))
        assert rows[0] == ("Name", "Value") and rows[1] == ("a", 1)

    res = client.delete(f"/api/notebooks/{notebook['id']}/artifacts/{art['id']}")
    assert res.json() == {"ok": True}
    assert client.get(f"/api/notebooks/{notebook['id']}/artifacts").json() == []
    if kind == "spreadsheet":
        assert not any(ARTIFACTS_DIR.glob(f"{notebook['id']}_*.xlsx"))


def test_artifact_unknown_kind_400(client, notebook):
    _seed(client, notebook["id"])
    assert client.post(f"/api/notebooks/{notebook['id']}/artifacts",
                       json={"kind": "hologram"}).status_code == 400


def test_artifact_empty_notebook_422(client, notebook):
    assert client.post(f"/api/notebooks/{notebook['id']}/artifacts",
                       json={"kind": "chart"}).status_code == 422


def test_notebook_delete_cleans_artifact_files(client, notebook):
    _seed(client, notebook["id"])
    art = client.post(f"/api/notebooks/{notebook['id']}/artifacts",
                      json={"kind": "spreadsheet"}).json()
    files = list(ARTIFACTS_DIR.glob(f"{notebook['id']}_*.xlsx"))
    assert files
    client.delete(f"/api/notebooks/{notebook['id']}")
    assert not any(Path(f).exists() for f in files)
    assert db.get_artifact(art["id"]) is None


# ---------- spec validation unit tests ----------

def test_validate_chart_spec_rejects_bad_shapes():
    with pytest.raises(ValueError):
        studio._validate_artifact_spec("chart", {"title": "t", "type": "scatter",
                                                 "labels": ["a"], "values": [1]})
    with pytest.raises(ValueError):
        studio._validate_artifact_spec("chart", {"title": "t", "type": "bar",
                                                 "labels": ["a", "b"], "values": [1]})
    with pytest.raises(studio.StudioError):
        studio._validate_artifact_spec("chart", {"error": "no numeric data"})


def test_mindgraph_generation_drops_dangling_links(client, notebook):
    """Links whose endpoints aren't real nodes can't be drawn, so they're pruned."""
    _seed(client, notebook["id"])
    spec = client.post(f"/api/notebooks/{notebook['id']}/artifacts",
                       json={"kind": "mindgraph"}).json()["spec"]
    assert spec["root"] == "Meridian Array"
    assert [b["label"] for b in spec["branches"]] == ["Output", "Site", "Cost"]
    assert spec["links"] == [{"from": "1.2 GW", "to": "$940M", "label": "drives"}]


def test_validate_mindgraph_rejects_bad_shapes():
    with pytest.raises(ValueError):   # no root
        studio._validate_artifact_spec("mindgraph", {
            "title": "t", "root": "  ",
            "branches": [{"label": "b", "children": ["c"]}]})
    with pytest.raises(ValueError):   # too few branches
        studio._validate_artifact_spec("mindgraph", {
            "title": "t", "root": "r", "branches": [{"label": "b", "children": ["c"]}]})
    with pytest.raises(ValueError):   # childless branch
        studio._validate_artifact_spec("mindgraph", {
            "title": "t", "root": "r",
            "branches": [{"label": "a", "children": ["c"]}, {"label": "b", "children": []}]})
    with pytest.raises(studio.StudioError):
        studio._validate_artifact_spec("mindgraph", {"error": "not enough material"})


def test_validate_mindgraph_caps_and_trims():
    spec = studio._validate_artifact_spec("mindgraph", {
        "title": "t", "root": " Root ",
        "branches": [
            {"label": " A ", "children": [f"c{i}" for i in range(9)]},
            {"label": "B", "children": ["x", "  ", "y"]},
        ],
        "links": [{"from": "c0", "to": "x", "label": "a much longer label than allowed"}],
    })
    assert spec["root"] == "Root" and spec["branches"][0]["label"] == "A"
    assert len(spec["branches"][0]["children"]) == 6      # capped
    assert spec["branches"][1]["children"] == ["x", "y"]  # blanks dropped
    assert len(spec["links"][0]["label"]) <= 20           # truncated


def test_validate_spreadsheet_pads_rows():
    spec = studio._validate_artifact_spec("spreadsheet", {
        "title": "t", "columns": ["a", "b", "c"], "rows": [["1"], ["1", "2", "3", "4"]],
    })
    assert spec["rows"] == [["1", "", ""], ["1", "2", "3"]]


# ---------- Phase 2: spreadsheet title sanitation ----------

@pytest.mark.parametrize("raw,expected", [
    ("Q3: Results", "Q3 Results"),
    ("Sales/Demand", "Sales Demand"),
    ("Data [final]", "Data final"),
    ("A\\B", "A B"),
    ("Why?", "Why"),
    ("*Metrics*", "Metrics"),
    ("[]", "Data"),
    ("   ", "Data"),
    ("'quoted'", "quoted"),
])
def test_safe_sheet_title_normalizes_illegal_characters(raw, expected):
    assert studio._safe_sheet_title(raw) == expected


def test_safe_sheet_title_enforces_excel_limit_and_unicode():
    out = studio._safe_sheet_title("Résumé financier — " + "é" * 40)
    assert len(out) <= 31
    assert out.startswith("Résumé financier")


def test_write_xlsx_accepts_illegal_and_unicode_titles(tmp_path):
    from openpyxl import load_workbook

    titles = ["Q3: Results", "Sales/Demand", "Data [final]", "A\\B",
              "Why?", "*Metrics*", "[]", "Résumé — données 2026"]
    for title in titles:
        path = tmp_path / "book.xlsx"
        studio.write_xlsx({"title": title, "columns": ["A"], "rows": [[1]]}, path)
        sheet = load_workbook(path).active.title
        assert sheet == studio._safe_sheet_title(title)
        assert sheet and len(sheet) <= 31


# ---------- Phase 2: retry boundaries ----------

def test_extract_json_rejects_non_object():
    with pytest.raises(studio.ModelOutputError):
        studio._extract_json("[1, 2, 3]")
    with pytest.raises(studio.ModelOutputError):
        studio._extract_json("no json here")


def test_validate_artifact_spec_rejects_non_dict():
    with pytest.raises(studio.ModelOutputError):
        studio._validate_artifact_spec("chart", ["not", "a", "dict"])


def test_validate_spreadsheet_rejects_non_list_row():
    with pytest.raises(studio.ModelOutputError):
        studio._validate_artifact_spec("spreadsheet", {
            "title": "t", "columns": ["a", "b"], "rows": ["oops"]})


def test_validate_mindgraph_ignores_non_scalar_children():
    spec = studio._validate_artifact_spec("mindgraph", {
        "title": "t", "root": "r",
        "branches": [
            {"label": "A", "children": ["good", {"bad": 1}, ["also", "bad"], 7]},
            {"label": "B", "children": ["x"]},
        ]})
    assert spec["branches"][0]["children"] == ["good", "7"]
    assert spec["branches"][1]["children"] == ["x"]


def test_generate_script_retries_and_ignores_broken_lines(monkeypatch):
    class LLM:
        def __init__(self):
            self.calls = 0

        def chat(self, messages, stream=False):
            self.calls += 1
            if self.calls == 1:
                # empty + missing speakers previously raised IndexError -> 500
                return '{"title":"t","lines":[{"speaker":"","text":"hi"},{"text":"x"}]}'
            return json.dumps({"title": "Good", "lines": [
                {"speaker": "A", "text": "one"}, {"speaker": "B", "text": "two"},
                {"speaker": "A", "text": "three"}, {"speaker": "B", "text": "four"}]})

    llm = LLM()
    monkeypatch.setattr(studio, "_gather_context", lambda nb, budget: "ctx")
    monkeypatch.setattr(studio, "get_llm", lambda: llm)
    out = studio.generate_script("nb")
    assert out["title"] == "Good" and len(out["lines"]) == 4
    assert llm.calls == 2


def _artifact_llm(replies):
    class LLM:
        def __init__(self):
            self.calls = 0

        def chat(self, messages, stream=False):
            reply = replies[min(self.calls, len(replies) - 1)]
            self.calls += 1
            return reply
    return LLM()


@pytest.mark.parametrize("first", [
    "[1, 2, 3]",                                                    # top-level array
    '{"title":"t","type":"scatter","labels":["a"],"values":[1]}',   # invalid chart type
    '{"type":"bar","labels":["a","b"],"values":[1]}',              # missing title
    '{"title":"t","type":"bar","labels":["a","b"],"values":["x","y"]}',  # non-numeric
])
def test_generate_artifact_spec_retries_malformed_structures(monkeypatch, first):
    good = json.dumps({"title": "OK", "type": "bar",
                       "labels": ["a", "b"], "values": [1, 2]})
    llm = _artifact_llm([first, good])
    monkeypatch.setattr(studio, "_gather_context", lambda nb, budget: "ctx")
    monkeypatch.setattr(studio, "get_llm", lambda: llm)
    spec = studio.generate_artifact_spec("nb", "chart")
    assert spec["title"] == "OK" and spec["values"] == [1.0, 2.0]
    assert llm.calls == 2


def test_generate_artifact_spec_gives_up_with_studio_error(monkeypatch):
    llm = _artifact_llm(["[1,2,3]"])
    monkeypatch.setattr(studio, "_gather_context", lambda nb, budget: "ctx")
    monkeypatch.setattr(studio, "get_llm", lambda: llm)
    with pytest.raises(studio.StudioError):
        studio.generate_artifact_spec("nb", "chart")
    assert llm.calls == 2


def test_generate_artifact_spec_does_not_retry_model_refusal(monkeypatch):
    llm = _artifact_llm([json.dumps({"error": "not enough numeric data"})])
    monkeypatch.setattr(studio, "_gather_context", lambda nb, budget: "ctx")
    monkeypatch.setattr(studio, "get_llm", lambda: llm)
    with pytest.raises(studio.StudioError):
        studio.generate_artifact_spec("nb", "chart")
    assert llm.calls == 1          # a model refusal is final, not retried


# ---------- Phase 2: infographic stat count ----------

def test_validate_infographic_keeps_all_six_stats():
    spec = studio._validate_artifact_spec("infographic", {
        "title": "t",
        "stats": [{"value": str(i), "label": f"s{i}"} for i in range(1, 7)],
        "sections": [{"heading": "h", "points": ["p"]}],
    })
    assert len(spec["stats"]) == 6


def test_validate_infographic_rejects_seven_stats():
    with pytest.raises(studio.ModelOutputError):
        studio._validate_artifact_spec("infographic", {
            "title": "t",
            "stats": [{"value": str(i), "label": f"s{i}"} for i in range(7)],
            "sections": [{"heading": "h", "points": ["p"]}]})


# ---------- Phase 3: explicit context budgets ----------

def test_studio_generators_pass_explicit_context_budgets(monkeypatch):
    from app.config import ARTIFACT_CONTEXT_CHARS, PODCAST_CONTEXT_CHARS

    seen = []
    monkeypatch.setattr(studio, "_gather_context",
                        lambda nb, budget: seen.append(budget) or "ctx")

    script_llm = _artifact_llm([json.dumps({"title": "t", "lines": [
        {"speaker": "A", "text": "1"}, {"speaker": "B", "text": "2"},
        {"speaker": "A", "text": "3"}, {"speaker": "B", "text": "4"}]})])
    monkeypatch.setattr(studio, "get_llm", lambda: script_llm)
    studio.generate_script("nb")
    assert seen[-1] == PODCAST_CONTEXT_CHARS

    chart_llm = _artifact_llm([json.dumps({
        "title": "c", "type": "bar", "labels": ["a", "b"], "values": [1, 2]})])
    monkeypatch.setattr(studio, "get_llm", lambda: chart_llm)
    studio.generate_artifact_spec("nb", "chart")
    assert seen[-1] == ARTIFACT_CONTEXT_CHARS


# ---------- Phase 8: grounded text artifacts ----------

TEXT_SPECS = {
    "briefing": {"title": "Brief", "sections": [{"heading": "H", "body": "Body."}]},
    "study_guide": {"title": "Study", "objectives": ["o"],
                    "key_concepts": [{"term": "t", "definition": "d"}],
                    "questions": [{"q": "q", "a": "a"}]},
    "faq": {"title": "FAQ", "items": [{"question": "q", "answer": "a"}]},
    "timeline": {"title": "Time", "events": [{"date": "Q1", "event": "e"}]},
    "source_summary": {"title": "Summary", "summary": "s", "key_points": ["p"]},
}


def test_all_text_artifact_kinds_are_registered():
    for kind in studio.TEXT_ARTIFACT_KINDS:
        assert kind in studio.ARTIFACT_PROMPTS


@pytest.mark.parametrize("kind,spec", TEXT_SPECS.items())
def test_text_artifact_generation(monkeypatch, kind, spec):
    llm = _artifact_llm([json.dumps(spec)])
    monkeypatch.setattr(studio, "_gather_context", lambda nb, budget: "ctx")
    monkeypatch.setattr(studio, "get_llm", lambda: llm)
    out = studio.generate_artifact_spec("nb", kind)
    assert out["title"] == spec["title"]
    assert llm.calls == 1


@pytest.mark.parametrize("kind,bad", [
    ("briefing", {"title": "B", "sections": []}),
    ("briefing", {"title": "B", "sections": [{"heading": "H"}]}),
    ("study_guide", {"title": "S", "objectives": [],
                     "key_concepts": [{"term": "t", "definition": "d"}],
                     "questions": [{"q": "q", "a": "a"}]}),
    ("faq", {"title": "F", "items": [{"question": "q"}]}),
    ("timeline", {"title": "T", "events": ["not an object"]}),
    ("source_summary", {"title": "Sum", "summary": "  ", "key_points": ["p"]}),
])
def test_text_artifact_validation_rejects_bad_shapes(kind, bad):
    with pytest.raises(studio.ModelOutputError):
        studio._validate_artifact_spec(kind, bad)
