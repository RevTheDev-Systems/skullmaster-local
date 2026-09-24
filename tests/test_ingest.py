"""Ingestion matrix: positive and negative fixtures for every supported type.

Fixtures are generated on the fly (no binaries committed) so each source type
has both a valid sample and a deliberately broken sample, and the documented
matrix is asserted against the parser.
"""

import sys
import types
import urllib.error

import fitz
import pytest
from docx import Document as DocxDocument
from openpyxl import Workbook

from app import ingest
from app.chunker import chunk_segments


@pytest.fixture(autouse=True)
def _no_real_youtube_network(monkeypatch):
    """YouTube tests must never reach the network: stub the oEmbed title lookup
    and block the audio downloader. Tests that exercise either patch it again."""

    def _blocked(url, video_id, max_minutes):
        raise AssertionError("test attempted a real YouTube audio download")

    monkeypatch.setattr(ingest, "_youtube_download_audio", _blocked)
    monkeypatch.setattr(ingest, "_youtube_title", lambda url, vid: f"YouTube video {vid}")


# The genuine downloader, kept before the autouse guard replaces it, so the few
# tests that exercise its internals can call it directly with a fake yt-dlp.
_REAL_DOWNLOAD = ingest._youtube_download_audio


# ---------- fixture builders ----------


def _txt(path, text="Hello text content."):
    path.write_text(text, encoding="utf-8")
    return path


def _pdf(path, pages=("Hello from page one.", "Page two has more content.")):
    doc = fitz.open()
    for text in pages:
        doc.new_page().insert_text((72, 72), text)
    doc.save(str(path))
    doc.close()
    return path


def _blank_pdf(path):
    doc = fitz.open()
    doc.new_page()  # page with no text layer (i.e. a scan)
    doc.save(str(path))
    doc.close()
    return path


def _docx(path, text="Hello DOCX world."):
    doc = DocxDocument()
    doc.add_heading("Title", level=1)
    doc.add_paragraph(text)
    doc.save(str(path))
    return path


def _xlsx(path):
    wb = Workbook()
    ws = wb.active
    ws.title = "Data"
    ws.append(["Name", "Value"])
    ws.append(["alpha", 1])
    ws.append(["beta", 2])
    wb.create_sheet("Empty")
    wb.save(str(path))
    return path


class _FakeSTT:
    def __init__(self, segments=None, error=None):
        self.segments = segments or []
        self.error = error

    def transcribe(self, path):
        if self.error:
            raise self.error
        return self.segments


class _FakeResponse:
    def __init__(self, body: bytes):
        self._body = body

    def read(self, n):
        return self._body[:n]

    def __enter__(self):
        return self

    def __exit__(self, *exc):
        return False


# ---------- matrix ----------


def test_ingestion_matrix_covers_required_kinds():
    kinds = {row["kind"] for row in ingest.INGESTION_MATRIX}
    assert {"pdf", "docx", "sheet", "text", "html", "url", "audio", "video"} <= kinds
    for row in ingest.INGESTION_MATRIX:
        assert row["citation"] in ("page", "passage", "seek")
        assert row["retry"] is True
    assert "ocr" not in kinds  # OCR is a future capability, not default


@pytest.mark.parametrize(
    "kind,ext,builder,expected_pages",
    [
        ("pdf", ".pdf", _pdf, 2),
        ("docx", ".docx", _docx, None),
        ("sheet", ".xlsx", _xlsx, None),
        ("text", ".txt", _txt, None),
        ("text", ".md", _txt, None),
        (
            "html",
            ".html",
            lambda p: _txt(p, "<html><body><h1>H</h1><p>Body words here.</p></body></html>"),
            None,
        ),
    ],
)
def test_supported_source_positives(tmp_path, kind, ext, builder, expected_pages):
    path = builder(tmp_path / f"sample{ext}")
    parsed_kind, pages, segments = ingest.parse_file(path)
    assert parsed_kind == kind
    assert pages == expected_pages
    assert segments
    assert any(text.strip() for _, text in segments)


# ---------- negatives: broken/empty/unsupported ----------


def test_zero_byte_and_corrupt_files_raise_ingest_error(tmp_path):
    cases = {
        "empty.txt": b"",
        "empty.pdf": b"",
        "corrupt.pdf": b"%PDF-1.4 this is not really a pdf",
        "corrupt.docx": b"PK\x03\x04 not actually a docx",
        "corrupt.xlsx": b"this is not a zip archive",
    }
    for name, data in cases.items():
        path = tmp_path / name
        path.write_bytes(data)
        with pytest.raises(ingest.IngestError):
            ingest.parse_file(path)


def test_unsupported_extension_raises(tmp_path):
    with pytest.raises(ingest.IngestError, match="Unsupported file type"):
        ingest.parse_file(_txt(tmp_path / "malware.exe", "MZ"))


def test_image_only_pdf_requires_ocr(tmp_path):
    with pytest.raises(ingest.IngestError, match="scanned"):
        ingest.parse_pdf(_blank_pdf(tmp_path / "scan.pdf"))


# ---------- OCR (optional engine, mocked) ----------


def test_scanned_pdf_is_ocrd_when_engine_available(tmp_path, monkeypatch):
    called = []
    monkeypatch.setattr(ingest, "ocr_configured", lambda: True)
    monkeypatch.setattr(ingest, "ocr_available", lambda: True)
    monkeypatch.setattr(
        ingest, "ocr_image", lambda png, lang: called.append(lang) or "Recovered scan text."
    )
    segments = ingest.parse_pdf(_blank_pdf(tmp_path / "scan.pdf"))
    assert called
    assert "Recovered scan text." in segments[0][1]


def test_text_pdf_is_never_ocrd(tmp_path, monkeypatch):
    called = []
    monkeypatch.setattr(ingest, "ocr_available", lambda: True)
    monkeypatch.setattr(ingest, "ocr_image", lambda png, lang: called.append(1) or "should not run")
    ingest.parse_pdf(_pdf(tmp_path / "text.pdf"))
    assert called == []


def test_ocr_runs_only_on_empty_pages(tmp_path, monkeypatch):
    called = []
    monkeypatch.setattr(ingest, "ocr_configured", lambda: True)
    monkeypatch.setattr(ingest, "ocr_available", lambda: True)
    monkeypatch.setattr(
        ingest, "ocr_image", lambda png, lang: called.append(1) or "page two recovered"
    )

    path = tmp_path / "mixed.pdf"
    doc = fitz.open()
    native_text = "Page one has real text. " * 10
    doc.new_page().insert_text((72, 72), native_text)
    doc.new_page()  # blank page (no text layer)
    doc.save(str(path))
    doc.close()

    segments = ingest.parse_pdf(path)
    assert len(called) == 1
    text = "\n".join(t for _, t in segments)
    assert "Page one has real text." in text and "page two recovered" in text


def test_scanned_pdf_rejected_when_ocr_disabled(tmp_path, monkeypatch):
    monkeypatch.setattr(ingest, "ocr_configured", lambda: False)
    with pytest.raises(ingest.IngestError, match="scanned"):
        ingest.parse_pdf(_blank_pdf(tmp_path / "scan.pdf"))


def test_empty_spreadsheet_raises(tmp_path):
    wb = Workbook()
    path = tmp_path / "empty.xlsx"
    wb.save(str(path))
    with pytest.raises(ingest.IngestError, match="no data"):
        ingest.parse_file(path)


# ---------- HTML parsing ----------


def test_html_is_parsed_not_raw(tmp_path):
    path = _txt(
        tmp_path / "p.html",
        "<html><body><script>evil()</script><h1>Head</h1><p>Body words.</p></body></html>",
    )
    kind, _, segments = ingest.parse_file(path)
    text = segments[0][1]
    assert kind == "html"
    assert "Body words." in text
    assert "<" not in text and "evil" not in text


# ---------- Unicode ----------


def test_unicode_content_is_preserved(tmp_path):
    text = "Données — 日本語 — Ω — satellite 🛰️"
    kind, _, segments = ingest.parse_file(_txt(tmp_path / "uni.txt", text))
    assert kind == "text"
    assert "日本語" in segments[0][1] and "🛰️" in segments[0][1]

    _, _, docx_segments = ingest.parse_file(_docx(tmp_path / "uni.docx", "Résumé — données"))
    assert "Résumé" in docx_segments[0][1]


def test_unusual_spreadsheet_values(tmp_path):
    wb = Workbook()
    ws = wb.active
    ws.title = "Wéird"
    ws.append(["A", None, "C"])
    ws.append([1, 2, 3])
    ws.merge_cells("A3:B3")
    ws["A3"] = "merged"
    ws.append([])
    path = tmp_path / "odd.xlsx"
    wb.save(str(path))

    _, _, segments = ingest.parse_file(path)
    text = segments[0][1]
    assert "Wéird" in text and "merged" in text and "C" in text


# ---------- large inputs ----------


def test_chunk_segments_honours_size_and_overlap():
    text = " ".join(f"Sentence {i} about arrays and sensors." for i in range(300))
    default = chunk_segments([(None, text)])
    small = chunk_segments([(None, text)], chunk_chars=800, overlap=100)
    assert len(small) > len(default)
    # a chunk is at most size + overlap (the carried tail) plus separators
    assert all(len(c["text"]) <= 1000 for c in small)
    assert all(c["page"] is None for c in small)


def test_large_document_chunks_without_error(tmp_path):
    big = "\n\n".join(f"Paragraph {i} " + "lorem ipsum dolor " * 20 for i in range(600))
    _, _, segments = ingest.parse_file(_txt(tmp_path / "big.txt", big))
    chunks = chunk_segments(segments)
    assert len(chunks) > 20
    assert all(c["text"].strip() for c in chunks)


# ---------- media (STT mocked) ----------


def test_media_timestamps_and_packing(tmp_path, monkeypatch):
    long_text = "sensor array telemetry " * 250  # > CHUNK_CHARS
    monkeypatch.setattr(
        ingest,
        "get_stt",
        lambda: _FakeSTT(
            [
                {"start": 0.0, "end": 5.0, "text": long_text},
                {"start": 42, "end": 50.0, "text": long_text},
            ]
        ),
    )
    kind, pages, segments = ingest.parse_file(_txt(tmp_path / "clip.mp4", "x"))
    assert kind == "video" and pages is None
    assert [sec for sec, _ in segments] == [0, 42]  # start seconds preserved


def test_media_no_speech_raises(tmp_path, monkeypatch):
    monkeypatch.setattr(ingest, "get_stt", lambda: _FakeSTT([]))
    with pytest.raises(ingest.IngestError, match="No speech"):
        ingest.parse_file(_txt(tmp_path / "silent.wav", "x"))


def test_media_transcription_failure_is_wrapped(tmp_path, monkeypatch):
    monkeypatch.setattr(ingest, "get_stt", lambda: _FakeSTT(error=RuntimeError("whisper boom")))
    with pytest.raises(ingest.IngestError, match="Transcription failed"):
        ingest.parse_file(_txt(tmp_path / "bad.mp3", "x"))


# ---------- URL ----------


def test_url_rejects_non_http_schemes():
    for url in ("file:///etc/passwd", "ftp://example.com/x", "javascript:alert(1)"):
        with pytest.raises(ingest.IngestError):
            ingest.parse_url(url)


def test_url_timeout_and_redirect_loop_are_wrapped(monkeypatch):
    def timeout(*a, **k):
        raise TimeoutError("timed out")

    monkeypatch.setattr(ingest.urllib.request, "urlopen", timeout)
    with pytest.raises(ingest.IngestError, match="Could not fetch"):
        ingest.parse_url("https://example.com/a")

    def redirect_loop(*a, **k):
        raise urllib.error.HTTPError("https://x", 302, "too many redirects", None, None)

    monkeypatch.setattr(ingest.urllib.request, "urlopen", redirect_loop)
    with pytest.raises(ingest.IngestError, match="Could not fetch"):
        ingest.parse_url("https://example.com/loop")


def test_url_oversize_body_rejected(monkeypatch):
    big = b"x" * (ingest.URL_MAX_BYTES + 10)
    monkeypatch.setattr(ingest.urllib.request, "urlopen", lambda *a, **k: _FakeResponse(big))
    with pytest.raises(ingest.IngestError, match="exceed"):
        ingest.parse_url("https://example.com/huge")


def test_url_without_readable_content_raises(monkeypatch):
    monkeypatch.setattr(
        ingest.urllib.request,
        "urlopen",
        lambda *a, **k: _FakeResponse(b"<html><body></body></html>"),
    )
    with pytest.raises(ingest.IngestError, match="No readable"):
        ingest.parse_url("https://example.com/empty")


def test_url_positive(monkeypatch):
    html = (
        "<html><head><title>My Article</title></head><body>"
        + "".join(f"<p>Paragraph {i} about the Atlas sensor network.</p>" for i in range(20))
        + "</body></html>"
    ).encode()
    monkeypatch.setattr(ingest.urllib.request, "urlopen", lambda *a, **k: _FakeResponse(html))
    title, segments = ingest.parse_url("https://example.com/article")
    assert title == "My Article"
    assert "Atlas sensor network" in segments[0][1]


# ---------- YouTube transcripts ----------


@pytest.mark.parametrize(
    "url,expected",
    [
        ("https://www.youtube.com/watch?v=aircAruvnKk", "aircAruvnKk"),
        ("https://youtube.com/watch?v=aircAruvnKk&t=42s", "aircAruvnKk"),
        ("https://m.youtube.com/watch?v=aircAruvnKk", "aircAruvnKk"),
        ("https://music.youtube.com/watch?v=aircAruvnKk", "aircAruvnKk"),
        ("https://youtu.be/aircAruvnKk", "aircAruvnKk"),
        ("https://youtu.be/aircAruvnKk?t=10", "aircAruvnKk"),
        ("https://www.youtube.com/shorts/aircAruvnKk", "aircAruvnKk"),
        ("https://www.youtube.com/embed/aircAruvnKk", "aircAruvnKk"),
        ("https://www.youtube.com/live/aircAruvnKk", "aircAruvnKk"),
    ],
)
def test_youtube_video_id_recognised(url, expected):
    assert ingest.youtube_video_id(url) == expected
    assert ingest.is_youtube_url(url) is True


@pytest.mark.parametrize(
    "url",
    [
        "https://example.com/watch?v=aircAruvnKk",
        "https://www.youtube.com/watch?v=tooshort",
        "https://www.youtube.com/",
        "https://notyoutube.com/aircAruvnKk",
        "not a url",
    ],
)
def test_youtube_video_id_rejects_non_videos(url):
    assert ingest.youtube_video_id(url) is None
    assert ingest.is_youtube_url(url) is False


class _Snippet:
    def __init__(self, start, text):
        self.start = start
        self.text = text


class _FakeTranscriptApi:
    """Stands in for youtube_transcript_api.YouTubeTranscriptApi (v1.x shape)."""

    snippets: list = []
    error: Exception | None = None
    calls: list = []

    def fetch(self, video_id):
        type(self).calls.append(video_id)
        if type(self).error:
            raise type(self).error
        return list(type(self).snippets)


def _patch_transcript_api(monkeypatch, snippets=None, error=None):
    import youtube_transcript_api

    _FakeTranscriptApi.snippets = snippets or []
    _FakeTranscriptApi.error = error
    _FakeTranscriptApi.calls = []
    monkeypatch.setattr(youtube_transcript_api, "YouTubeTranscriptApi", _FakeTranscriptApi)
    return _FakeTranscriptApi


def test_parse_youtube_returns_timestamped_transcript(monkeypatch):
    _patch_transcript_api(
        monkeypatch,
        snippets=[
            _Snippet(4.2, "This is a three."),
            _Snippet(9.9, "It is sloppily written."),
            _Snippet(700.0, "A later point entirely."),
        ],
    )
    monkeypatch.setattr(ingest, "_youtube_title", lambda url, vid: "A Video — A Channel")
    title, segments = ingest.parse_youtube("https://www.youtube.com/watch?v=aircAruvnKk")
    assert title == "A Video — A Channel"
    # A >45s gap starts a new block, so the transcript is timestamped, not one blob.
    assert [sec for sec, _ in segments] == [4, 700]
    assert "This is a three." in segments[0][1]
    assert "A later point" in segments[1][1]


def test_parse_youtube_groups_dense_snippets_into_one_block(monkeypatch):
    _patch_transcript_api(monkeypatch, snippets=[_Snippet(i * 2, f"line {i}") for i in range(10)])
    monkeypatch.setattr(ingest, "_youtube_title", lambda url, vid: "T")
    _, segments = ingest.parse_youtube("https://youtu.be/aircAruvnKk")
    assert len(segments) == 1 and segments[0][0] == 0


def test_parse_youtube_without_captions_fails_clearly(monkeypatch):
    _patch_transcript_api(monkeypatch, error=RuntimeError("no transcript found"))
    monkeypatch.setattr(ingest, "youtube_fallback_configured", lambda: False)
    with pytest.raises(ingest.IngestError, match="captions"):
        ingest.parse_youtube("https://www.youtube.com/watch?v=aircAruvnKk")


def test_parse_youtube_empty_transcript_fails(monkeypatch):
    _patch_transcript_api(monkeypatch, snippets=[])
    monkeypatch.setattr(ingest, "youtube_fallback_configured", lambda: False)
    with pytest.raises(ingest.IngestError, match="empty transcript"):
        ingest.parse_youtube("https://www.youtube.com/watch?v=aircAruvnKk")


def test_youtube_captions_win_and_skip_the_download(monkeypatch):
    """Captions are preferred: a captioned video must never download audio."""
    _patch_transcript_api(monkeypatch, snippets=[_Snippet(0.0, "Captions are here.")])
    monkeypatch.setattr(ingest, "_youtube_title", lambda url, vid: "T")
    monkeypatch.setattr(ingest, "youtube_fallback_configured", lambda: True)
    _, segments = ingest.parse_youtube("https://youtu.be/aircAruvnKk")
    assert segments == [(0, "Captions are here.")]


def test_youtube_without_captions_falls_back_to_local_transcription(monkeypatch, tmp_path):
    """No captions → download the audio, transcribe it locally, keep timestamps."""
    _patch_transcript_api(monkeypatch, error=RuntimeError("no transcript found"))
    monkeypatch.setattr(ingest, "_youtube_title", lambda url, vid: "Captioned-less — Chan")
    monkeypatch.setattr(ingest, "youtube_fallback_configured", lambda: True)
    workdir = tmp_path / "dl"
    workdir.mkdir()
    audio = workdir / "aircAruvnKk.webm"
    audio.write_bytes(b"fake audio bytes")
    monkeypatch.setattr(ingest, "_youtube_download_audio", lambda *a: audio)
    long_line = "spoken words " * 300  # exceeds CHUNK_CHARS, forcing a second block
    monkeypatch.setattr(
        ingest,
        "get_stt",
        lambda: types.SimpleNamespace(
            transcribe=lambda path: [
                {"start": 0.0, "end": 2.0, "text": long_line.strip()},
                {"start": 90.0, "end": 92.0, "text": "A later sentence."},
            ]
        ),
    )
    title, segments = ingest.parse_youtube("https://www.youtube.com/watch?v=aircAruvnKk")
    assert title == "Captioned-less — Chan"
    # local transcription keeps the timestamps, exactly like uploaded media
    assert [sec for sec, _ in segments] == [0, 90]
    assert "spoken words" in segments[0][1]
    assert "A later sentence." in segments[1][1]
    # the temporary download is removed; only the transcript survives
    assert not workdir.exists()


def test_youtube_fallback_can_be_disabled(monkeypatch):
    _patch_transcript_api(monkeypatch, error=RuntimeError("no transcript found"))
    monkeypatch.setattr(ingest, "youtube_fallback_configured", lambda: False)
    with pytest.raises(ingest.IngestError, match="fallback disabled"):
        ingest.parse_youtube("https://youtu.be/aircAruvnKk")


def test_youtube_fallback_without_speech_fails(monkeypatch, tmp_path):
    _patch_transcript_api(monkeypatch, error=RuntimeError("no transcript found"))
    monkeypatch.setattr(ingest, "youtube_fallback_configured", lambda: True)
    workdir = tmp_path / "dl"
    workdir.mkdir()
    monkeypatch.setattr(ingest, "_youtube_download_audio", lambda *a: workdir / "a.webm")
    monkeypatch.setattr(
        ingest, "get_stt", lambda: types.SimpleNamespace(transcribe=lambda path: [])
    )
    with pytest.raises(ingest.IngestError, match="No speech detected"):
        ingest.parse_youtube("https://youtu.be/aircAruvnKk")
    assert not workdir.exists()  # temp download cleaned up even on failure


def test_youtube_audio_duration_is_capped_before_downloading(monkeypatch):
    """A long video is refused before any bytes are fetched."""
    downloaded = []

    class _FakeYDL:
        def __init__(self, options):
            self.options = options

        def __enter__(self):
            return self

        def __exit__(self, *exc):
            return False

        def extract_info(self, url, download=False):
            return {"duration": 3 * 60 * 60}

        def download(self, urls):
            downloaded.append(urls)

    monkeypatch.setattr(ingest, "_youtube_download_audio", _REAL_DOWNLOAD)
    monkeypatch.setitem(sys.modules, "yt_dlp", types.SimpleNamespace(YoutubeDL=_FakeYDL))
    with pytest.raises(ingest.IngestError, match="capped at 120 minutes"):
        ingest._youtube_download_audio("https://youtu.be/aircAruvnKk", "aircAruvnKk", 120)
    assert not downloaded


def test_youtube_audio_download_returns_the_audio_file(monkeypatch):
    seen = {}

    class _FakeYDL:
        def __init__(self, options):
            self.options = options

        def __enter__(self):
            return self

        def __exit__(self, *exc):
            return False

        def extract_info(self, url, download=False):
            return {"duration": 60}

        def download(self, urls):
            seen["urls"] = urls
            seen["options"] = self.options
            target = ingest.Path(self.options["outtmpl"].replace("%(id)s", "aircAruvnKk"))
            target.with_suffix(".webm").write_bytes(b"audio")
            # a leftover .part file must be ignored when choosing the result
            target.with_suffix(".webm.part").write_bytes(b"partial")

    monkeypatch.setattr(ingest, "_youtube_download_audio", _REAL_DOWNLOAD)
    monkeypatch.setitem(sys.modules, "yt_dlp", types.SimpleNamespace(YoutubeDL=_FakeYDL))
    path = ingest._youtube_download_audio("https://youtu.be/aircAruvnKk", "aircAruvnKk", 120)
    assert path.name == "aircAruvnKk.webm" and path.read_bytes() == b"audio"
    assert seen["urls"] == ["https://youtu.be/aircAruvnKk"]
    assert seen["options"]["noplaylist"] is True  # never expand a playlist link


def test_youtube_audio_download_reports_missing_yt_dlp(monkeypatch):
    monkeypatch.setattr(ingest, "_youtube_download_audio", _REAL_DOWNLOAD)
    monkeypatch.setitem(sys.modules, "yt_dlp", None)  # import yt_dlp -> ImportError
    with pytest.raises(ingest.IngestError, match="yt-dlp"):
        ingest._youtube_download_audio("https://youtu.be/aircAruvnKk", "aircAruvnKk", 120)


def test_parse_url_routes_youtube_to_transcripts(monkeypatch):
    monkeypatch.setattr(ingest, "parse_youtube", lambda url: ("Video", [(0, "spoken words")]))
    title, segments = ingest.parse_url("https://youtu.be/aircAruvnKk")
    assert title == "Video" and segments == [(0, "spoken words")]


def test_parse_url_rejects_boilerplate_only_pages(monkeypatch):
    """A JS-rendered page (e.g. a video site) yields only chrome — refuse it
    rather than ingesting junk that the model could then cite."""
    chrome = (
        "<html><body><nav>About Press Copyright Contact us Creators "
        "Advertise Developers Terms Privacy</nav></body></html>"
    ).encode()
    monkeypatch.setattr(ingest.urllib.request, "urlopen", lambda *a, **k: _FakeResponse(chrome))
    with pytest.raises(ingest.IngestError, match="readable article content"):
        ingest.parse_url("https://example.com/video-page")


def test_ingestion_matrix_includes_youtube():
    rows = {row["kind"]: row for row in ingest.INGESTION_MATRIX}
    assert rows["youtube"]["citation"] == "seek"
    assert "timestamp" in rows["youtube"]["metadata"]


def test_ingestion_matrix_has_no_duplicate_or_missing_kinds():
    """The matrix is the contract; a duplicate (which silently replaces a kind
    in a dict) or a dropped kind must fail loudly."""
    kinds = [row["kind"] for row in ingest.INGESTION_MATRIX]
    assert len(kinds) == len(set(kinds)), f"duplicate kinds: {kinds}"
    assert set(kinds) >= {
        "pdf",
        "docx",
        "sheet",
        "text",
        "html",
        "url",
        "youtube",
        "audio",
        "video",
        "image",
    }
